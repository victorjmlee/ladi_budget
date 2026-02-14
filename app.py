#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
웹 UI 제공 + 엑셀 추출 API (Python 파이프라인).
실행: python app.py  →  http://localhost:3580
"""
import io
import os
import tempfile
from pathlib import Path

from flask import Flask, request, send_file, send_from_directory

# 프로젝트 루트(kolo)에서 import
from pdf_extract import extract_text_by_page, extract_by_codes, extract_by_keywords
from excel_build import build_excel_aoa

app = Flask(__name__, static_folder=None)
ROOT = Path(__file__).resolve().parent
PORT = int(os.environ.get("PORT", 3580))


def parse_keywords(raw: str):
    return [s.strip() for s in raw.replace(",", "\n").splitlines() if s.strip()]


@app.route("/")
def index():
    return send_from_directory(ROOT, "index.html")


@app.route("/<path:path>")
def static_file(path):
    if path == "index.html":
        return send_from_directory(ROOT, path)
    if Path(ROOT / path).is_file():
        return send_from_directory(ROOT, path)
    return "", 404


@app.route("/api/extract-excel", methods=["POST"])
def extract_excel():
    if "pdf" not in request.files:
        return {"error": "PDF 파일이 없습니다."}, 400
    file = request.files["pdf"]
    if not file or not file.filename or not file.filename.lower().endswith(".pdf"):
        return {"error": "PDF 파일을 선택하세요."}, 400
    raw_codes = (request.form.get("codes") or "").strip()
    raw_keywords = (request.form.get("keywords") or "").strip()
    codes = [s for s in parse_keywords(raw_codes) if len(s) == 3 and s.isdigit()]
    kw_list = [s for s in parse_keywords(raw_keywords) if s]
    # 하위호환: codes 필드 없이 keywords에 3자리 숫자만 온 경우 코드로 처리
    if not codes and kw_list and all(len(k) == 3 and k.isdigit() for k in kw_list):
        codes = kw_list
        kw_list = []
    if not codes and not kw_list:
        return {"error": "코드(201, 207) 또는 키워드를 입력하세요."}, 400

    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            file.save(tmp.name)
            tmp_path = tmp.name
        try:
            page_texts = extract_text_by_page(tmp_path)
            code_results = extract_by_codes(page_texts, codes) if codes else []
            kw_results = extract_by_keywords(page_texts, kw_list) if kw_list else []
            if not code_results and not kw_results:
                return {"error": "매칭되는 줄이 없습니다."}, 400
            code_aoa = build_excel_aoa(code_results, page_texts) if code_results else None
            kw_aoa = build_excel_aoa(kw_results, page_texts) if kw_results else None
        finally:
            try:
                os.unlink(tmp_path)
            except Exception:
                pass

        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from parser import EXCEL_HEADER

        # 단일 시트에 코드/키워드 결과 합치기 + 검색방식 컬럼(H)
        def _norm_key(s):
            return (s or "").replace(" ", "").replace("ㅇ", "○").strip()

        combined = [EXCEL_HEADER + ["검색방식"]]
        code_keys = set()

        if code_aoa:
            for row in code_aoa[1:]:
                combined.append(row + ["코드"])
                key = _norm_key(row[3])
                if key and key != "편성목":
                    code_keys.add(key)

        if kw_aoa:
            for row in kw_aoa[1:]:
                key = _norm_key(row[3])
                if key and key != "편성목" and key in code_keys:
                    for existing in combined[1:]:
                        if _norm_key(existing[3]) == key and existing[7] == "코드":
                            existing[7] = "코드+키워드"
                            break
                else:
                    combined.append(row + ["키워드"])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "예산추출"
        header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        kw_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        both_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        for r, row in enumerate(combined, 1):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=val or "")
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                elif len(row) > 7:
                    stype = row[7]
                    if stype == "코드+키워드":
                        cell.fill = both_fill
                    elif stype == "키워드":
                        cell.fill = kw_fill

        ws.auto_filter.ref = f"A1:H{len(combined)}"

        for col_idx in range(1, 9):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, len(combined) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = cell.value
                if val is not None:
                    s = str(val)
                    length = sum(2 if "\uAC00" <= ch <= "\uD7A3" or ord(ch) > 127 else 1 for ch in s)
                    max_len = max(max_len, length)
            ws.column_dimensions[col_letter].width = min(max(max_len / 2 + 2, 8), 55)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        base_name = Path(file.filename).stem or "검색결과"
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=base_name + ".xlsx",
        )
    except Exception as e:
        return {"error": str(e)}, 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=PORT, debug=False)
